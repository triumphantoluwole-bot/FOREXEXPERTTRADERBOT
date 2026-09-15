# bot.py
"""
FOREX EXPERT TRADER BOT — @FOREXEXPERTTRADERBOT
Extracts 7 key fields from PDF / DOCX / TXT documents.
"""

import os
import logging
import tempfile

import fitz  # PyMuPDF
from docx import Document

from telegram import Update
from telegram.ext import (
    Application,
    CommandHandler,
    MessageHandler,
    filters,
    ContextTypes,
)

from extractor import TenderExtractor, ExtractedData


logging.basicConfig(
    format="%(asctime)s - %(name)s - %(levelname)s - %(message)s",
    level=logging.INFO,
)
logger = logging.getLogger(__name__)


ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}
TELEGRAM_MAX_LEN = 4000


# ---------------------------------------------------------------- Handlers

async def start(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "🚀 *FOREX EXPERT TRADER BOT*\n\n"
        "I extract key data from your documents.\n\n"
        "Send me a *PDF*, *DOCX*, or *TXT* file and I will reply with:\n"
        "• Requirements\n"
        "• Submission deadline\n"
        "• Required documents\n"
        "• Eligibility criteria\n"
        "• Evaluation criteria\n"
        "• Budget information\n"
        "• Important clauses",
        parse_mode="Markdown",
    )


async def help_command(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Simply send a PDF, DOCX, or TXT file and I'll extract the key data.",
    )


async def handle_text(update: Update, context: ContextTypes.DEFAULT_TYPE):
    await update.message.reply_text(
        "Please send a PDF, DOCX, or TXT file to extract data.",
    )


# ------------------------------------------------------- Text extraction

def _extract_pdf_text(path: str) -> str:
    with fitz.open(path) as pdf:
        return "\n".join(page.get_text() for page in pdf)


def _extract_docx_text(path: str) -> str:
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text.strip():
                    parts.append(cell.text)
    return "\n".join(parts)


def _extract_txt_text(path: str) -> str:
    with open(path, "r", encoding="utf-8", errors="ignore") as f:
        return f.read()


# --------------------------------------------------------- Formatting

def _format_response(data: ExtractedData) -> str:
    sections = [
        ("📋 *Requirements*", data.requirements),
        ("📅 *Submission Deadline*", data.submission_deadline),
        ("📁 *Required Documents*", data.required_documents),
        ("✅ *Eligibility Criteria*", data.eligibility_criteria),
        ("📊 *Evaluation Criteria*", data.evaluation_criteria),
        ("💰 *Budget Information*", data.budget_information),
        ("⚠️ *Important Clauses*", data.important_clauses),
    ]

    out = "*📑 EXTRACTION RESULTS*\n\n"
    for title, content in sections:
        out += f"{title}\n"
        if isinstance(content, list):
            if content:
                for item in content:
                    out += f"• {item}\n"
            else:
                out += "• Not found\n"
        else:
            out += f"{content}\n"
        out += "\n"
    return out


# ----------------------------------------------------- Document handler

async def handle_document(update: Update, context: ContextTypes.DEFAULT_TYPE):
    message = update.message
    doc = message.document

    if not doc:
        await message.reply_text("❌ Please send a document file.")
        return

    file_name = doc.file_name or "document"
    ext = os.path.splitext(file_name)[1].lower()

    if ext not in ALLOWED_EXTENSIONS:
        await message.reply_text(
            "❌ Unsupported file type. Please send a PDF, DOCX, or TXT file.",
        )
        return

    await message.reply_text(
        f"📄 Processing `{file_name}`... Please wait.",
        parse_mode="Markdown",
    )

    tmp_path = None
    try:
        file = await context.bot.get_file(doc.file_id)
        with tempfile.NamedTemporaryFile(delete=False, suffix=ext) as tmp:
            tmp_path = tmp.name
        await file.download_to_drive(tmp_path)

        if ext == ".pdf":
            text = _extract_pdf_text(tmp_path)
        elif ext == ".docx":
            text = _extract_docx_text(tmp_path)
        else:
            text = _extract_txt_text(tmp_path)

        if not text.strip():
            await message.reply_text(
                "❌ No readable text found. The file may be a scanned image.",
            )
            return

        extractor = TenderExtractor(text)
        data = extractor.extract_all()
        response = _format_response(data)

        if len(response) <= TELEGRAM_MAX_LEN:
            await message.reply_text(response, parse_mode="Markdown")
        else:
            for i in range(0, len(response), TELEGRAM_MAX_LEN):
                await message.reply_text(
                    response[i:i + TELEGRAM_MAX_LEN],
                    parse_mode="Markdown",
                )

    except Exception as e:
        logger.exception("Error processing document: %s", e)
        await message.reply_text(
            "❌ An error occurred while processing the document.",
        )
    finally:
        if tmp_path and os.path.exists(tmp_path):
            try:
                os.unlink(tmp_path)
            except OSError:
                pass


# --------------------------------------------------------------- Main

def main():
    token = os.environ.get("TELEGRAM_BOT_TOKEN")
    if not token:
        raise ValueError("TELEGRAM_BOT_TOKEN environment variable not set.")

    application = Application.builder().token(token).build()

    application.add_handler(CommandHandler("start", start))
    application.add_handler(CommandHandler("help", help_command))
    application.add_handler(
        MessageHandler(filters.Document.ALL, handle_document)
    )
    application.add_handler(
        MessageHandler(filters.TEXT & ~filters.COMMAND, handle_text)
    )

    logger.info("Bot starting...")
    application.run_polling(allowed_updates=Update.ALL_TYPES)


if __name__ == "__main__":
    main()
